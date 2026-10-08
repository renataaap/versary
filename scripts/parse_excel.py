"""Read Excel without pandas/native numerical DLLs; preserve the import JSON contract."""
import json
import sys
from datetime import datetime, date
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

def unique_headers(values):
    columns, used = [], set()
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else f"Unnamed: {index}"
        base = base or f"coluna_{index+1}"
        name, suffix = base, 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        columns.append(name)
    return columns

def read_sheets(file_path):
    if file_path.suffix.lower() in {".xls", ".xld"}:
        import xlrd
        workbook = xlrd.open_workbook(str(file_path))
        try:
            for sheet in workbook.sheets():
                values = []
                for row in range(sheet.nrows):
                    items = []
                    for cell in sheet.row(row):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
                            value = None
                        items.append(value)
                    values.append(items)
                yield sheet.name, values
        finally:
            workbook.release_resources()
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

def parse_workbook(file_path):
    if not file_path.is_file():
        raise FileNotFoundError("O arquivo Excel temporário não foi encontrado.")
    sheets = []
    for name, values in read_sheets(file_path):
        if not values:
            continue
        # Ignore trailing empty columns from Excel formatting, not genuine headers.
        width = max((index+1 for row in values for index, value in enumerate(row) if value is not None), default=0)
        if not width:
            continue
        columns = unique_headers(list(values[0])[:width])
        rows = []
        for row in values[1:]:
            if not any(value is not None and value != "" for value in row):
                continue
            rows.append({column: row[index] if index < len(row) and row[index] is not None else "" for index, column in enumerate(columns)})
        sheets.append({"sheet":name,"columns":columns,"rows":rows,"totalRows":len(rows)})
    if not sheets:
        raise ValueError("Nenhuma aba contém colunas para importar.")
    return {"sheets":sheets}

def main():
    if len(sys.argv) != 2:
        raise ValueError("Informe o caminho do arquivo Excel.")
    print(json.dumps(parse_workbook(Path(sys.argv[1])),ensure_ascii=False,default=str,allow_nan=False))

if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error":str(error)},ensure_ascii=False))
        raise SystemExit(1)